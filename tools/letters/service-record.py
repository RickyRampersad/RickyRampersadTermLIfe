#!/usr/bin/env python3
"""Add each client's own service record to the Transition Send list: what the branch team has done for them.

Asked for on 24 September 2026: "we have had many touchpoints and were able to resolve; they are to see us as
from onboarding and a team service, as we have all the data on service levels." The letters print it in a panel
under the facts, "Handled for you by our branch team", one cell per count, and a cell whose count is blank or
zero is cut out like a blank fact. A client with nothing on record gets the plain line about the team instead.

  python3 service-record.py <send-list.csv> <out.csv> --portfolio <Branch Portfolio copy .csv>
          --tasks <Task.json> --records <CLIENT_PORTFOLIO__c.json> --logbook <Log Book workbook .xlsx>

Adds four columns, the letters' merge fields:
  svc_docs       documents on the client's policies in the Log Book: checked by the document team, sent on
  svc_requests   completed Salesforce tasks of service work for them (not the automated notes, not group admin)
  svc_reminders  premium reminders sent to them
  svc_birthday   the month of the last birthday note, "March 2026"

The counts are quoted to the client, so only the precise links count:
  - a task on the client's own portfolio record, or naming one of their policy numbers
  - a task on a contact, only where the contact's name is the client's own (surname and a given name), so a
    spouse's or a child's record is never lent to them
  - never an Account link: a household or employer account would lend one client another's history
  - the Log Book by policy number, one line per policy, document type and day
Every service task in the branch's Salesforce on 24 September 2026 was created by the branch team (sales
support, the branch manager's assistant, the personal assistants, the branch manager), none by an agent, which
is what lets the letter say "our branch team". Salesforce tasks begin in 2025 and the Log Book in 2023, so a
count understates a long relationship and never overstates one.

Client data never enters the repository: run this outside it, on copies, and keep what it writes there.
"""
import argparse, collections, csv, datetime as dt, json, re

from openpyxl import load_workbook

AUTOMATED = r'birthday|special occasion|anniversary|premium due|important reminder|delivery update|next steps|agent assigned|' \
            r'courtesy reminder|health insurance renewal|experience with our agents|we want to hear'
# group administration, internal reports and team notices: work, but not work done for this client
INTERNAL = r'weekly task update|t-\s?(health|pensions|life) group|business opportunity|funds in disbursement|orphan|premium due listing|' \
           r'daily digest|aim\W*digest|dashboard|recruit|resume\b|attendance|licen[cs]e|minutes of|training|registered for|' \
           r'new case email notification|case status|^ms - |^online - |^email: ms - |exam registration'


def records(path):
    o = json.load(open(path, encoding='utf-8'))
    return o if isinstance(o, list) else next(v for v in o.values() if isinstance(v, list))


def pn(v):
    return re.sub(r'\D', '', re.sub(r'\.0$', '', str(v or '').strip())).lstrip('0')


def toks(s):
    return [t for t in re.findall(r'[a-z]+', (s or '').lower()) if len(t) > 1]


def day(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', str(v or ''))
    return dt.date(*map(int, m.groups())) if m else None


def lbkind(doc):
    d = doc.upper()
    for k, pat in (('beneficiary', r'BENEF|NOMINAT'), ('surrender', r'SURREND'), ('reinstatement', r'REINST'), ('refund', r'REFUND'),
                   ('loan', r'\bLOAN'), ('cancellation', r'CANCEL|FREE.?LOOK|NOT TAKEN'),
                   ('payment change', r'ACH|DIRECT DEBIT|BANKERS|SALARY DED|FUTURE PREMIUM|BILLING'),
                   ('contract', r'ACKNOWLEDG|POLICY CONTRACT|DELIVERY'), ('claim', r'CLAIM|DISCHARGE'),
                   ('application', r'APPLICATION|SIGNED SCRIPT|OAS|PROOF OF ID|MEDICAL')):
        if re.search(pat, d): return k
    return 'other'


def logbook(path):
    """(policy, day, document type) for every line of the three tabs that carry a policy number."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out = set()
    it = wb['Form responses - Digital Log'].iter_rows(values_only=True); h = next(it)
    ix = {str(k).strip(): i for i, k in enumerate(h) if k}
    for r in it:
        txt = str(r[ix['Document Submission']] or '') if ix['Document Submission'] < len(r) else ''
        when = day(r[0])
        for line in txt.split('\n'):
            m = re.search(r'Policy\s*#\s*:\s*([^,]+),\s*Document\s*:\s*([^,]+)', line, re.I)
            if m and pn(m.group(1)) and when:
                out.add((pn(m.group(1)), when, lbkind(m.group(2))))
    for name, pc, dc, dtc in (('BU', 'Policy', 'Document', 1), ('MAIL PREP', 'Policy #', 'Document', 0)):
        it = wb[name].iter_rows(values_only=True); h = next(it); ix = {}
        for i, k in enumerate(h):
            if k: ix.setdefault(str(k).strip(), i)
        for r in it:
            if ix[pc] >= len(r): continue
            p, doc, when = pn(r[ix[pc]]), str(r[ix[dc]] or '').strip(), day(r[dtc])
            if p and doc and when:
                out.add((p, when, lbkind(doc)))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('send_list'); ap.add_argument('out')
    ap.add_argument('--portfolio', required=True); ap.add_argument('--tasks', required=True)
    ap.add_argument('--records', required=True); ap.add_argument('--logbook', required=True)
    ap.add_argument('--today', default=dt.date.today().isoformat())
    a = ap.parse_args()
    today = dt.date.fromisoformat(a.today)

    pol2cn, cname = {}, {}
    for r in csv.DictReader(open(a.portfolio, newline='', encoding='utf-8-sig')):
        cn = r['Client Number'].strip()
        pol2cn[pn(r['Number'])] = cn
        cname.setdefault(cn, r['Client'])
    cp2cn, con2cn = {}, collections.defaultdict(set)
    for r in records(a.records):
        cn = str(r.get('Client_Number__c') or '').strip()
        if cn not in cname:
            cn = next((pol2cn[pn(r.get(f))] for f in ('POLICY__c', 'Policy_Number__c', 'Policy_ID__c') if pn(r.get(f)) in pol2cn), '')
        if cn:
            cp2cn[r['Id']] = cn
            if r.get('Contact__c'): con2cn[r['Contact__c']].add(cn)

    def same_person(who, cn):
        x, y = toks(who), toks(cname.get(cn, ''))
        return bool(x and y) and x[-1] in y and bool(set(x[:-1]) & set(y))

    reqs, rems, bdays = (collections.defaultdict(set) for _ in range(3))
    for t in records(a.tasks):
        subj = t.get('Subject') or ''
        d = day(t.get('ActivityDate') or t.get('CreatedDate'))
        if not d or d > today or re.search(INTERNAL, subj, re.I): continue
        who = (t.get('Who') or {}).get('Name') or ''
        cns = {cp2cn[t['WhatId']]} if t.get('WhatId') in cp2cn else set()
        cns |= {pol2cn[pn(m.group(0))] for m in re.finditer(r'\b\d{8,11}\b', subj) if pn(m.group(0)) in pol2cn}
        cns |= {cn for cn in con2cn.get(t.get('WhoId') or '', ()) if len(con2cn[t['WhoId']]) <= 4 and same_person(who, cn)}
        for cn in cns:
            if re.search(r'birthday|special occasion', subj, re.I): bdays[cn].add(d)
            elif re.search(r'premium due|important reminder', subj, re.I): rems[cn].add((d, subj))
            elif not re.search(AUTOMATED, subj, re.I) and t.get('Status') == 'Completed': reqs[cn].add((d, subj.lower()))
    docs = collections.defaultdict(set)
    for p, d, k in logbook(a.logbook):
        if p in pol2cn and d <= today: docs[pol2cn[p]].add((p, d, k))

    rows = list(csv.DictReader(open(a.send_list, newline='', encoding='utf-8-sig')))
    cols = list(rows[0].keys()) if rows else []
    for c in ('svc_docs', 'svc_requests', 'svc_reminders', 'svc_birthday'):
        if c not in cols: cols.append(c)
    filled = collections.Counter()
    for r in rows:
        cn = (r.get('Client number') or '').strip()
        r['svc_docs'] = str(len(docs[cn]) or '')
        r['svc_requests'] = str(len(reqs[cn]) or '')
        r['svc_reminders'] = str(len(rems[cn]) or '')
        r['svc_birthday'] = max(bdays[cn]).strftime('%B %Y') if bdays[cn] else ''
        for c in ('svc_docs', 'svc_requests', 'svc_reminders', 'svc_birthday'):
            filled[c] += bool(r[c])
        filled['any'] += any(r[c] for c in ('svc_docs', 'svc_requests', 'svc_reminders', 'svc_birthday'))
    with open(a.out, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(rows)
    print(f'{len(rows)} rows written to {a.out}; with a service record: {filled["any"]} '
          f'(documents {filled["svc_docs"]}, requests {filled["svc_requests"]}, reminders {filled["svc_reminders"]}, '
          f'birthday {filled["svc_birthday"]})')


if __name__ == '__main__':
    main()
